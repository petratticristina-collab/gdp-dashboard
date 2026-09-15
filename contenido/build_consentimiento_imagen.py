"""Genera el consentimiento de imagen, voz y testimonio (DOCX editable) para pacientes."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "consentimiento-imagen-testimonio-metodo-petratti.docx")
LOGO = os.path.join(HERE, "logo-metodo-petratti.png")
TURQ = RGBColor(0x00, 0xA8, 0xA4)
GRIS = RGBColor(0x3F, 0x45, 0x48)

doc = Document()
for s in doc.sections:
    s.top_margin, s.bottom_margin = Cm(2), Cm(2)
    s.left_margin, s.right_margin = Cm(2.2), Cm(2.2)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def h(text, size=12, color=TURQ, space=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def para(text, size=10.5, italic=False, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.font.color.rgb = GRIS
    return p


def check(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("☐  " + text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = GRIS


def field(label, width_lines=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + " ")
    r.bold = True
    r.font.color.rgb = GRIS
    p.add_run("_" * 60 if width_lines == 1 else "_" * 90)


# ---------- Cabecera ----------
hdr = doc.sections[0].header
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
if os.path.exists(LOGO):
    hp.add_run().add_picture(LOGO, width=Cm(4.2))

ftr = doc.sections[0].footer.paragraphs[0]
ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = ftr.add_run("Dra. Cristina Petratti · Método Petratti · Consentimiento de imagen, voz y testimonio · v1 · 2026")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(0x88, 0x88, 0x8C)

# ---------- Título ----------
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = t.add_run("CONSENTIMIENTO INFORMADO PARA LA CAPTACIÓN, USO Y DIFUSIÓN\nDE IMAGEN, VOZ Y TESTIMONIO")
tr.bold = True
tr.font.size = Pt(14)
tr.font.color.rgb = TURQ
para("Ley Orgánica 1/1982, de protección civil del derecho al honor, a la intimidad personal y familiar y a la "
     "propia imagen · Reglamento (UE) 2016/679 (RGPD), arts. 6.1.a, 7 y 9.2.a · Ley Orgánica 3/2018 (LOPDGDD)",
     size=8.5, italic=True, after=10)

# ---------- 1. Partes ----------
h("1. Responsable del tratamiento")
para("Dra. Cristina Petratti · Médica de familia, especialista en obesidad · Nº de colegiada: ____________ · "
     "NIF: ____________ · Domicilio profesional: ____________________________________________ · "
     "Correo de contacto para el ejercicio de derechos: ____________________________ · "
     "Delegado/a de protección de datos (si procede): ____________________________")

h("2. Persona que otorga el consentimiento")
field("Nombre y apellidos:")
field("DNI / NIE / Pasaporte:")
field("Teléfono:")
field("Correo electrónico:")
para("La persona firmante es mayor de edad y otorga este consentimiento de forma libre, específica, informada e "
     "inequívoca. Este consentimiento es independiente de la relación asistencial: negarse a firmarlo o revocarlo "
     "no afecta en nada a la atención médica recibida.", after=6)

# ---------- 3. Objeto ----------
h("3. Objeto")
para("Autorizo a la Dra. Cristina Petratti a captar, fijar, reproducir y difundir mi imagen, mi voz y mi "
     "testimonio sobre mi experiencia personal en relación con mi salud y mi proceso de tratamiento, en los "
     "términos, canales y con los límites que marco a continuación.")
para("Descripción del contenido concreto (fecha de grabación, formato, duración aproximada, tema):", after=2)
para("______________________________________________________________________________________________")
para("______________________________________________________________________________________________", after=6)

# ---------- 4. Qué autorizo ----------
h("4. Elementos que autorizo (marcar solo los que procedan)")
check("Imagen fija (fotografía)")
check("Imagen en movimiento (vídeo)")
check("Voz (audio)")
check("Mi nombre de pila")
check("Mi nombre y apellidos")
check("Mi edad aproximada")
check("Mención a mi diagnóstico o categoría de tratamiento (p. ej. «tratamiento para la obesidad»), sin "
      "nombres de fármacos, dosis, cifras de peso ni resultados analíticos")
check("Uso de mi imagen con el rostro visible (si no se marca, se pixelará o se usará plano sin rostro)")

# ---------- 5. Canales ----------
h("5. Canales y finalidad (marcar los que procedan)")
para("Finalidad: divulgación en salud y comunicación de la actividad profesional de la Dra. Petratti. El contenido "
     "no se asociará a promesas de resultados, cifras de pérdida de peso ni a la promoción de medicamentos.")
check("Instagram (@crispetratti) y Threads")
check("YouTube (@dra.cristinapetratti)")
check("TikTok")
check("Página web y boletín por correo electrónico")
check("Presentaciones docentes, congresos y formación a profesionales")
check("Medios de comunicación (prensa, radio, televisión) previa comunicación a la persona firmante")
para("Si el contenido lo publica la propia persona en sus redes sociales:", after=2)
check("Autorizo a la Dra. Petratti a compartir en sus canales (stories, repost, cita) el contenido que yo publique "
      "voluntariamente en mis propias redes sobre mi experiencia, sin modificarlo")
check("Autorizo a la Dra. Petratti a responder públicamente a comentarios o preguntas sobre mi caso, únicamente "
      "sobre lo que yo haya hecho público, sin añadir datos clínicos")
check("Declaro que no recibo ninguna contraprestación (descuento, sesiones, regalo) por publicar en mis redes. "
      "Si la recibiera, identificaré la publicación como colaboración (#publi)")
para("Ámbito territorial: internacional, por la naturaleza de internet. Duración de la autorización: "
     "☐ 2 años · ☐ 5 años · ☐ indefinida, desde la fecha de firma; en todos los casos revocable.", after=6)

# ---------- 6. Garantías ----------
h("6. Garantías que asume la Dra. Petratti")
for g in [
    "Podré revisar el contenido final antes de su publicación y pedir cambios o retirarme, sin necesidad de "
    "justificarlo.  ☐ Sí, quiero revisarlo   ☐ No es necesario",
    "No se editará mi testimonio de forma que altere su sentido.",
    "No se difundirán datos clínicos concretos (analíticas, kilos, dosis, nombres de fármacos) salvo que yo lo "
    "autorice expresamente por escrito para un contenido concreto.",
    "El contenido se publicará con la indicación de que es una experiencia personal y no sustituye la "
    "valoración médica individual.",
    "Mis datos identificativos no se cederán a terceros, salvo a las plataformas necesarias para la publicación, "
    "que actúan con sus propias condiciones de servicio.",
]:
    check(g) if g.startswith("Podré") else para("• " + g, after=2)

# ---------- 7. Contraprestación ----------
h("7. Contraprestación")
para("☐ Esta cesión es gratuita.   ☐ Se acuerda una contraprestación de ____________ (describir). "
     "En ningún caso la contraprestación se condiciona al contenido del testimonio ni a resultados clínicos.")

# ---------- 8. Revocación y derechos ----------
h("8. Revocación y ejercicio de derechos")
para("Puedo revocar este consentimiento en cualquier momento, sin justificación, mediante escrito al correo "
     "indicado en el apartado 1. La Dra. Petratti retirará el contenido de los canales bajo su control en un plazo "
     "máximo de 15 días desde la recepción. La revocación no tiene efectos retroactivos sobre difusiones ya "
     "realizadas ni sobre copias efectuadas por terceros ajenos a su control.")
para("Puedo ejercer los derechos de acceso, rectificación, supresión, limitación, oposición y portabilidad ante el "
     "responsable, y presentar una reclamación ante la Agencia Española de Protección de Datos (www.aepd.es). Los "
     "datos se conservarán mientras el contenido esté publicado y, después, bloqueados durante los plazos de "
     "prescripción legal.", after=6)

# ---------- 9. Declaración ----------
h("9. Declaración y firma")
para("Declaro que he leído y comprendido este documento, que he podido hacer preguntas y que han sido resueltas, "
     "y que otorgo mi consentimiento de forma libre. Recibo una copia firmada.", after=10)

tbl = doc.add_table(rows=2, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
cells = tbl.rows[0].cells
cells[0].text = "Firma de la persona que consiente\n\n\n\n_____________________________\nNombre:\nFecha y lugar:"
cells[1].text = "Firma de la Dra. Cristina Petratti\n\n\n\n_____________________________\nNº colegiada:\nFecha y lugar:"
for c in cells:
    for p in c.paragraphs:
        for r in p.runs:
            r.font.size = Pt(10)
            r.font.color.rgb = GRIS

doc.add_paragraph()
para("Anexo opcional · Autorización específica de datos clínicos para un contenido concreto: autorizo la mención "
     "de ______________________________________ en el contenido descrito en el apartado 3.  "
     "Firma: __________________  Fecha: __________", size=9, after=0)

doc.save(OUT)
print("OK", OUT)
